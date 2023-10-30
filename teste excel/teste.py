from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl import load_workbook

lista_musicas = load_workbook("teste excel/Lista_de_músicas1.xlsx")
ws = lista_musicas.active
print(ws.tables.items())