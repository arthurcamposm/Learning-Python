from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
import pandas as pd


lista_musicas = load_workbook("teste excel/Lista_de_músicas1 copy.xlsx")
aba_ativa = lista_musicas.active
tabela = aba_ativa.tables['Tabela1']
"""read_lista = pd.read_excel('teste excel/Lista_de_músicas1 copy.xlsx')
cont = 0
maior_num = 0
for num in read_lista['Número']:
    cont += 1
    if cont == 1:
        maior_num = num
    if num > maior_num:
        maior_num = num
print(maior_num)"""
música = str(input('Digite o nome da música: ')).capitalize().strip()
versão = str(input('Digite a versão da música: ')).capitalize().strip()
categoria = str(input('Digite a categoria da música entre Adoração, Celebração, Comunhão, Páscoa, Missões, Guerra e Reflexão: ')).capitalize().strip()
drive = str(input('A música está no drive? [Sim / Não] ')).capitalize().strip()
cifra = str(input('Foi montada uma cifra para a música? [Sim / Não] ')).capitalize().strip()
letra = 'Não'
mp3 = ''
youtube = str(input('A música está no YouTube? [Sim / Não] ')).capitalize().strip()
arranjo = ''
if drive == 'Sim':
    linkdrive = str(input('Adicione o link do drive: '))
else:
    linkdrive = ''
status = str(input('Adicione o Status da música entre Relembrar, Relembrando, Tirada ou Rever: ')).capitalize().strip()
df_emp = pd.DataFrame ([música, versão, categoria, drive, cifra, letra, mp3, youtube, arranjo, linkdrive, status])
with pd.ExcelWriter('Lista_de_músicas1 copy.xlsx', mode="a", if_sheet_exists="overlay") as writer:
    df_emp.to_excel(writer, startrow=139, header=False, index=False)


